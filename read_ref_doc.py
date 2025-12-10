from docx import Document
import os

file_path = "Greendeks Autofill Rehber Doküman.docx"

if not os.path.exists(file_path):
    print(f"Error: File not found at {file_path}")
else:
    doc = Document(file_path)
    print("--- PARAGRAPHS ---")
    for para in doc.paragraphs:
        if para.text.strip():
            print(para.text)
    
    print("\n--- TABLES ---")
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            print(" | ".join(row_text))
        print("-" * 20)
